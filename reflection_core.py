"""Input validation and document assembly for the reflection website."""
import io
from datetime import timedelta
from pathlib import Path

import docx
from docxcompose.composer import Composer
from docxtpl import DocxTemplate
from google.genai import types
from jinja2 import Environment, StrictUndefined
from pythainlp.tokenize import word_tokenize
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


def source_part(name, content):
    if not content:
        raise ValueError("ไฟล์โครงการสอนว่างเปล่า กรุณาเลือกไฟล์ใหม่")
    if len(content) > 15 * 1024 * 1024:
        raise ValueError("ไฟล์ต้องมีขนาดไม่เกิน 15 MB")
    suffix = Path(name).suffix.lower()
    if suffix == ".docx":
        document = docx.Document(io.BytesIO(content))
        lines = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        text = "\n".join(lines)
        if not text.strip():
            raise ValueError("ไม่พบข้อความใน Word หากเป็นภาพสแกน กรุณาอัปโหลด PDF หรือรูปภาพแทน")
        return types.Part.from_text(text=text)
    if suffix == ".doc":
        return types.Part.from_bytes(data=content, mime_type="application/msword")
    if suffix == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp874")
        return types.Part.from_text(text=text)
    mime = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}
    if suffix not in mime:
        raise ValueError("ไม่รองรับไฟล์ชนิดนี้")
    return types.Part.from_bytes(data=content, mime_type=mime[suffix])


def parse_holidays(text, count):
    holidays = {}
    for line in text.splitlines():
        if not line.strip():
            continue
        number, separator, reason = line.partition(":")
        if not separator or not number.strip().isdigit() or not reason.strip():
            raise ValueError("กรอกวันงดสอนในรูปแบบ สัปดาห์:เหตุผล หนึ่งรายการต่อบรรทัด")
        number = int(number.strip())
        if not 1 <= number <= count:
            raise ValueError(f"สัปดาห์ที่งดสอนต้องอยู่ระหว่าง 1–{count}")
        holidays[number] = reason.strip()
    return holidays


def validate_weeks(data, count, holidays):
    if not isinstance(data, dict) or not isinstance(data.get("weeks"), list):
        raise ValueError("AI ส่งข้อมูลไม่ตรงรูปแบบ กรุณาลองสร้างใหม่")
    weeks = data["weeks"]
    numbers = [w.get("week") for w in weeks if isinstance(w, dict)]
    if len(numbers) != count or any(type(n) is not int for n in numbers) or sorted(numbers) != list(range(1, count + 1)):
        raise ValueError(f"AI ส่งข้อมูลสัปดาห์ไม่ครบหรือซ้ำ ต้องมีสัปดาห์ที่ 1–{count} กรุณาลองสร้างใหม่")
    for week in weeks:
        number = week["week"]
        week["is_holiday"] = number in holidays
        week["off_reason"] = holidays.get(number, "-")
        if number in holidays:
            week.update(topic="งดการจัดการเรียนรู้", student_eval="งดสอน", teacher_eval="งดสอน", problem_solution=holidays[number])
        elif any(not isinstance(week.get(field), str) or not week[field].strip() for field in ("topic", "student_eval", "teacher_eval", "problem_solution")):
            raise ValueError(f"ข้อมูลสัปดาห์ที่ {number} ไม่ครบ กรุณาลองสร้างใหม่")
    return sorted(weeks, key=lambda w: w["week"])


def lesson_date(start, week, weekday):
    # The chosen start date is the first day of the first seven-day teaching cycle.
    return start + timedelta(weeks=week - 1, days=(weekday - start.weekday()) % 7)


def format_topic_for_form(topic, line_width=30, max_lines=4):
    """Keep the topic compact while retaining its beginning and key ending."""
    if not isinstance(topic, str):
        return topic
    lines = []
    for paragraph in topic.splitlines() or [""]:
        current = ""
        for token in word_tokenize(paragraph.strip(), engine="newmm", keep_whitespace=True):
            candidate = current + token
            if len(candidate) > line_width and current.strip():
                lines.append(current.rstrip())
                current = token.lstrip()
            else:
                current = candidate
        lines.append(current.rstrip())
    if len(lines) > max_lines:
        if max_lines == 1:
            lines = [lines[0].rstrip(" .") + "..."]
        else:
            lines = lines[:max_lines - 1] + [lines[-1].lstrip()]
            lines[-1] = "..." + lines[-1].rstrip(" .") + "..."
    return "\n".join(lines)


def _center_date_time_cells(document, context):
    values = {
        str(context.get("date", "")).strip(),
        str(context.get("date_display", "")).strip(),
        str(context.get("time", "")).strip(),
        str(context.get("time_display", "")).strip(),
    }
    values.discard("")
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() not in values:
                    continue
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def render_document(template, contexts):
    composer = None
    for context in contexts:
        tpl = DocxTemplate(io.BytesIO(template))
        tpl.render(context, jinja_env=Environment(undefined=StrictUndefined), autoescape=True)
        buffer = io.BytesIO()
        tpl.save(buffer)
        buffer.seek(0)
        document = docx.Document(buffer)
        _center_date_time_cells(document, context)
        if composer is None:
            composer = Composer(document)
        else:
            composer.doc.add_page_break()
            composer.append(document)
    if composer is None:
        raise ValueError("ไม่มีข้อมูลสำหรับสร้างเอกสาร")
    output = io.BytesIO()
    composer.save(output)
    return output.getvalue()
