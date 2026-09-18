import io
import unittest
from datetime import date

import docx
from streamlit.testing.v1 import AppTest
from reflection_core import (
    lesson_date, parse_holidays, validate_weeks, render_document,
    source_part, format_topic_for_form,
)


class ReflectionTests(unittest.TestCase):
    def test_calendar_never_precedes_start(self):
        start = date(2026, 9, 17)
        self.assertEqual(lesson_date(start, 1, 0), date(2026, 9, 21))
        self.assertEqual(lesson_date(start, 2, 0), date(2026, 9, 28))

    def test_reject_incomplete_ai_and_invalid_holiday(self):
        with self.assertRaises(ValueError):
            validate_weeks({'weeks': [{'week': 1}]}, 18, {})
        with self.assertRaises(ValueError):
            parse_holidays('19:หยุด', 18)
        self.assertEqual(parse_holidays('8:หยุด', 18), {8: 'หยุด'})

    def test_word_render_escape_merge_and_input(self):
        document = docx.Document()
        document.add_paragraph('{{ teacher_name }} / สัปดาห์ {{ week }}')
        buffer = io.BytesIO()
        document.save(buffer)
        output = render_document(buffer.getvalue(), [
            {'teacher_name': 'ครู A & B <C>', 'week': 1},
            {'teacher_name': 'ครู A & B <C>', 'week': 2},
        ])
        result = docx.Document(io.BytesIO(output))
        text = '\n'.join(p.text for p in result.paragraphs)
        self.assertIn('ครู A & B <C> / สัปดาห์ 1', text)
        self.assertIn('ครู A & B <C> / สัปดาห์ 2', text)
        self.assertIn('สัปดาห์ 2', source_part('course.docx', output).text)

    def test_word_render_centers_date_and_time_cells(self):
        document = docx.Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '{{ date_display }}'
        table.cell(0, 1).text = '{{ time_display }}'
        buffer = io.BytesIO()
        document.save(buffer)
        output = render_document(buffer.getvalue(), [{
            'date_display': 'วันจันทร์ 18 พฤษภาคม 2569',
            'time_display': 'เวลา 15.30-16.30 น.',
        }])
        result = docx.Document(io.BytesIO(output))
        cells = result.tables[0].rows[0].cells
        self.assertEqual(cells[0].vertical_alignment, docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER)
        self.assertEqual(cells[1].vertical_alignment, docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER)
        self.assertEqual(cells[0].paragraphs[0].alignment, docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(cells[1].paragraphs[0].alignment, docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER)

    def test_topic_is_limited_to_four_lines(self):
        topic = "หน่วยการเรียนรู้ ระบบเครือข่ายคอมพิวเตอร์และการติดตั้งอุปกรณ์สำหรับงานจริงในสถานศึกษา พร้อมการบำรุงรักษาและการแก้ไขปัญหาเชิงระบบ"
        formatted = format_topic_for_form(topic)
        self.assertLessEqual(len(formatted.splitlines()), 4)
        self.assertTrue(formatted.endswith("..."))

    def test_topic_wraps_at_thai_word_boundaries(self):
        formatted = format_topic_for_form("การจัดการเรียนรู้ระบบเครือข่ายคอมพิวเตอร์", line_width=15)
        self.assertNotIn("การจัดการเรีย\nนรู้", formatted)
        self.assertNotIn("เครือข่\nาย", formatted)

    def test_interface_auth_levels_and_validation(self):
        app = AppTest.from_file('app.py').run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        app.text_input[0].set_value('wrong')
        app.button[0].click().run()
        self.assertEqual(len(app.error), 1)
        app.session_state.authenticated = True
        app.run()
        self.assertEqual(len(app.exception), 0)
        next(s for s in app.selectbox if s.label == 'ระดับคุณวุฒิ:').select('ปวส.').run()
        self.assertEqual(len(app.metric), 0)
        next(b for b in app.button if 'สร้างร่าง' in b.label).click().run()
        self.assertEqual(len(app.warning), 1)
        self.assertEqual(len(app.exception), 0)


if __name__ == '__main__':
    unittest.main()
